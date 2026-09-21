import asyncio
import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import selectinload, sessionmaker

from app.models.apartment import Apartment
from app.models.resident import Resident

DATABASE_URL = "postgresql+asyncpg://postgres:postgres@localhost:5432/astana_opera_housing"


async def main():
    engine = create_async_engine(DATABASE_URL, echo=False)
    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async with async_session() as session:
        stmt = (
            select(Apartment)
            .where(Apartment.residential_complex_name == "Москва")
            .options(selectinload(Apartment.residents))
        )
        res = await session.execute(stmt)
        apts = res.scalars().all()

        print("Current DB Apartments in 'Москва':")
        for a in apts:
            res_names = [r.full_name for r in a.residents if r.is_active]
            print(f"  Apt ID {a.id}, Apt No: {a.apartment_number}, Resident: {res_names}, Subtype: {a.apartment_subtype}")

    # Inspect Excel file data/рассрочка и выкупленные общее поступление 2222.xlsx
    excel_path = os.path.join(os.path.dirname(__file__), "data", "рассрочка и выкупленные общее поступление 2222.xlsx")
    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        print("\nSheets in excel:", wb.sheetnames)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n--- Sheet: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                row_str = " | ".join([str(c) for c in row if c is not None])
                if any(name in row_str for name in ["Жагипаров", "Куанышбек", "Науанов", "Мукашев", "Усин", "123", "138", "160", "27", "63"]):
                    print(row_str[:250])

if __name__ == "__main__":
    asyncio.run(main())
