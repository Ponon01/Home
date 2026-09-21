import asyncio
import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import selectinload, sessionmaker

from app.models.apartment import Apartment
from app.models.resident import Resident

DATABASE_URL = "postgresql+asyncpg://postgres:postgres@localhost:5432/astana_opera_housing"


async def main():
    engine = create_async_engine(DATABASE_URL, echo=False)
    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async with async_session() as session:
        stmt = (
            select(Apartment)
            .where(Apartment.residential_complex_name == "Москва")
            .options(selectinload(Apartment.residents))
        )
        res = await session.execute(stmt)
        apts = res.scalars().all()

        print("=== Current DB Apartments in 'Москва' ===")
        for a in apts:
            residents = [f"{r.full_name} (id:{r.id})" for r in a.residents if r.is_active]
            print(f"Apt ID {a.id}, Apt No: {a.apartment_number}, Street: {a.street}, Address: {a.address}, Residents: {residents}")

    print("\n=== Searching Excel files ===")
    files = [
        "рассрочка и выкупленные общее поступление 2222.xlsx",
        "Книга1.xlsx",
    ]
    targets = ["Жагипаров", "Жағыпаров", "Куанышбек", "Қуанышбек", "Науанов", "Мукашев", "Мұқашев", "Усин", "Үсін"]

    for fname in files:
        fpath = os.path.join(os.path.dirname(__file__), "data", fname)
        if not os.path.exists(fpath):
            print(f"File not found: {fname}")
            continue
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                row_str = " | ".join([str(c) for c in row if c is not None])
                if any(t.lower() in row_str.lower() for t in targets):
                    print(f"[{fname} -> {sname} : L{r_idx}] {row_str[:220]}")

if __name__ == "__main__":
    asyncio.run(main())
