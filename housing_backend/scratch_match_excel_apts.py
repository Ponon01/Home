import asyncio
import sys
import os
import openpyxl
from collections import defaultdict
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from sqlalchemy import select
from app.models.apartment import Apartment

DATABASE_URL = "postgresql+asyncpg://postgres:postgres@localhost:5432/astana_opera_housing"

# Simple name normalization to match Excel complex to DB complex
def normalize_name(name):
    if not name:
        return ""
    name = name.lower().strip()
    # Replace common spelling variants
    name = name.replace(" ", "").replace("-", "").replace("—", "")
    if "лазурный" in name:
        return "лазурныйквартал"
    if "сапа" in name or "сити" in name:
        return "сапа2007"
    if "азирбаев" in name:
        return "улкаазирбаева"
    return name

async def main():
    engine = create_async_engine(DATABASE_URL, echo=False)
    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    
    # 1. Read Excel rows
    wb = openpyxl.load_workbook("data/Книга1.xlsx", data_only=True)
    sheet = wb.worksheets[0]
    
    excel_apts = defaultdict(list)
    for r in range(2, sheet.max_row + 1):
        idx = sheet.cell(row=r, column=2).value
        apt_num = sheet.cell(row=r, column=7).value
        full_addr = sheet.cell(row=r, column=8).value
        
        if not full_addr:
            continue
            
        # Try to parse complex name from full_addr or guess
        complex_guess = ""
        full_addr_lower = full_addr.lower()
        if "лазурный" in full_addr_lower:
            complex_guess = "Лазурный квартал"
        elif "зерде" in full_addr_lower:
            complex_guess = "Зерде"
        elif "сармат" in full_addr_lower:
            complex_guess = "Сармат"
        elif "хан-тенгри" in full_addr_lower or "хан тенгри" in full_addr_lower:
            complex_guess = "Хан-тенгри"
        elif "виктория" in full_addr_lower:
            complex_guess = "Виктория"
        elif "жагалау" in full_addr_lower:
            complex_guess = "Жагалау-3"
        elif "сапа" in full_addr_lower:
            complex_guess = "Сапа-2007"
        elif "акку" in full_addr_lower:
            complex_guess = "Акку"
        elif "общежитие" in full_addr_lower:
            complex_guess = "Общежитие"
        elif "азирбаев" in full_addr_lower:
            complex_guess = "ул. К. Азирбаева"
        elif "браво" in full_addr_lower:
            complex_guess = "Браво"
        elif "compass" in full_addr_lower:
            complex_guess = "Compass North"
        elif "respublika" in full_addr_lower or "республика" in full_addr_lower:
            # Distinguish from Общежитие на Республики 81
            if "республики 81" not in full_addr_lower:
                complex_guess = "Respublika"
            else:
                complex_guess = "Общежитие"
        elif "нур-сая" in full_addr_lower or "нур сая" in full_addr_lower or "нурсая" in full_addr_lower:
            complex_guess = "НУР-САЯ"
        elif "москва" in full_addr_lower:
            complex_guess = "Москва"
            
        if not complex_guess:
            # Fallback if no match
            continue
            
        amort = sheet.cell(row=r, column=47).value or 0
        deduct = sheet.cell(row=r, column=48).value or 0
        tax = sheet.cell(row=r, column=49).value or 0
        
        try:
            amort = float(amort)
        except:
            amort = 0.0
        try:
            deduct = float(deduct)
        except:
            deduct = 0.0
        try:
            tax = float(tax)
        except:
            tax = 0.0
            
        apt_str = str(apt_num).strip() if apt_num is not None else ""
        if not apt_str and "-" in full_addr:
            # try to parse from last part of address
            parts = full_addr.split("-")
            if parts:
                apt_str = parts[-1].strip()
                
        key = (normalize_name(complex_guess), apt_str)
        excel_apts[key].append({
            "complex": complex_guess,
            "apt": apt_str,
            "amortization": amort,
            "deduction": deduct,
            "tax": tax
        })
        
    print(f"Total parsed Excel apartment keys: {len(excel_apts)}")
    
    # 2. Query DB apartments
    async with async_session() as session:
        res = await session.execute(select(Apartment))
        db_apts = res.scalars().all()
        
        matched_count = 0
        unmatched_db = []
        for a in db_apts:
            c_norm = normalize_name(a.residential_complex_name)
            apt_str = str(a.apartment_number).strip() if a.apartment_number else ""
            key = (c_norm, apt_str)
            if key in excel_apts:
                matched_count += 1
            else:
                unmatched_db.append(f"{a.residential_complex_name} Apt {a.apartment_number}")
                
        print(f"Matched DB apartments: {matched_count} / {len(db_apts)}")
        print(f"Unmatched DB samples: {unmatched_db[:20]}")

if __name__ == "__main__":
    asyncio.run(main())
