import openpyxl
from pathlib import Path

def analyze_sheets():
    # 1. Rent
    path_rent = Path("data") / "1Расчет по квартирам 2026 (АРЕНДА).xlsx"
    wb_rent = openpyxl.load_workbook(path_rent, read_only=True, data_only=True)
    
    # 2. Purchase/Installment
    path_pur = Path("data") / "рассрочка и выкупленные общее поступление 2222.xlsx"
    wb_pur = openpyxl.load_workbook(path_pur, read_only=True, data_only=True)
    
    # Count rows and IIN presence in Rent Sheet '2026 Март'
    ws_rent = wb_rent['2026 Март']
    rent_rows_with_iin = 0
    rent_rows_without_iin = 0
    fios_with_iin = {}
    
    headers = [cell.value for cell in next(ws_rent.iter_rows(min_row=1, max_row=1))]
    fio_col = headers.index('ФИО') if 'ФИО' in headers else -1
    iin_col = headers.index('ИИН') if 'ИИН' in headers else -1
    
    print(f"Rent headers: {headers}")
    print(f"FIO column: {fio_col}, IIN column: {iin_col}")
    
    for row in ws_rent.iter_rows(min_row=2):
        row_vals = [cell.value for cell in row]
        if not any(row_vals):
            continue
        fio = str(row_vals[fio_col]).strip() if (fio_col != -1 and len(row_vals) > fio_col and row_vals[fio_col]) else None
        iin = str(row_vals[iin_col]).strip() if (iin_col != -1 and len(row_vals) > iin_col and row_vals[iin_col]) else None
        
        if fio:
            if iin:
                rent_rows_with_iin += 1
                fios_with_iin[fio] = iin
            else:
                rent_rows_without_iin += 1
                
    print(f"Rent: rows with FIO and IIN: {rent_rows_with_iin}, rows without IIN: {rent_rows_without_iin}")
    
    # Check purchase files
    for sheet_name in ['Лист1', 'Лист2']:
        ws_pur = wb_pur[sheet_name]
        pur_headers = [cell.value for cell in next(ws_pur.iter_rows(min_row=5 if sheet_name == 'Лист1' else 3, max_row=5 if sheet_name == 'Лист1' else 3))]
        print(f"\nPurchase sheet {sheet_name} headers: {pur_headers}")
        
        fio_idx = -1
        for idx, h in enumerate(pur_headers):
            if h and 'ФИО' in str(h).upper():
                fio_idx = idx
                break
        
        matched = 0
        unmatched = 0
        for r_idx in range((6 if sheet_name == 'Лист1' else 4), ws_pur.max_row + 1):
            row_vals = [ws_pur.cell(row=r_idx, column=c_idx+1).value for c_idx in range(len(pur_headers))]
            if not any(row_vals):
                continue
            fio = str(row_vals[fio_idx]).strip() if (fio_idx != -1 and len(row_vals) > fio_idx and row_vals[fio_idx]) else None
            if fio:
                if fio in fios_with_iin:
                    matched += 1
                else:
                    unmatched += 1
                    
        print(f"Purchase sheet {sheet_name}: FIO matches with Rent: {matched}, no match: {unmatched}")
        
    wb_rent.close()
    wb_pur.close()

if __name__ == "__main__":
    analyze_sheets()
