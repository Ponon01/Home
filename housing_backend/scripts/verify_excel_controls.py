# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

DATA = Path("data")


def find(*keys):
    for p in sorted(DATA.glob("*.xlsx")):
        if p.name.startswith("~$") or p.name.startswith("_"):
            continue
        low = p.name.casefold().replace("ё", "е")
        if all(k in low for k in keys):
            return p
    return None


def num(ws, r, c):
    v = ws.cell(r, c).value
    try:
        return float(v or 0)
    except Exception:
        return 0


# --- EARLY ---
p = find("доссроч", "выкуп") or find("досроч", "выкуп")
print("EARLY", p)
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
print("headers r3", [ws.cell(3, c).value for c in range(1, 12)])
init = bal = val = 0
n = 0
for r in range(4, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    if not addr or not fio or addr.lower().startswith("итого"):
        continue
    init += num(ws, r, 6)
    bal += num(ws, r, 9)
    val += num(ws, r, 10)
    n += 1
print("early n", n, "init", round(init, 2), "bal", round(bal, 2), "val", round(val, 2))
wb.close()

# --- INSTALLMENT ---
p = find("рассроч")
print("INST", p)
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb["Лист1"]
print("headers r5", [ws.cell(5, c).value for c in range(1, 14)])
init = pay = mon = 0
n = 0
for r in range(6, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    if not addr or not fio:
        continue
    low = addr.lower()
    if "итого" in low or "информац" in low:
        continue
    init += num(ws, r, 6)
    pay += num(ws, r, 11)
    mon += num(ws, r, 12)
    n += 1
print("inst n", n, "init", round(init, 2), "pay", round(pay, 2), "mon", round(mon, 2))
wb.close()

# --- SALES ---
p = find("сумма")
print("SALES", p)
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb.active
for r in range(1, (ws.max_row or 0) + 1):
    vals = [ws.cell(r, c).value for c in range(1, 14)]
    if any(v is not None and str(v).strip() not in ("", "None") for v in vals):
        print(r, vals)
wb.close()

# --- FUND ---
p = find("информац", "фонд")
print("FUND", p)
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb["Свод по ЖК русс"]
for r in (6, 7, 8, 24):
    print("fund", r, [ws.cell(r, c).value for c in range(1, 19)])
rent = guest = 0
for r in range(8, 24):
    name = str(ws.cell(r, 1).value or "")
    if not name or name.upper().startswith("ИТОГО"):
        continue
    def ni(c):
        v = ws.cell(r, c).value
        try:
            return int(v or 0)
        except Exception:
            return 0
    rent += ni(17)
    guest += ni(18)
    print(name[:40], "r", ni(17), "g", ni(18), "rem", ni(16), "tot", ni(3))
print("SUM rent", rent, "guest", guest)
wb.close()
