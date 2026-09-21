# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

DATA = Path("data")


def find(*keys):
    for p in sorted(DATA.glob("*.xlsx")):
        if p.name.startswith("~$") or p.name.startswith("_"):
            continue
        low = p.name.casefold().replace("ё", "е")
        if all(k in low for k in keys):
            return p
    return None


def num(ws, r, c):
    v = ws.cell(r, c).value
    try:
        return float(v or 0)
    except Exception:
        return 0


p = find("доссроч", "выкуп") or find("досроч", "выкуп")
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]

groups = {
    "early": ("досроч", "доссроч"),
    "immediate": ("сразу",),
    "installment": ("действующ", "рассроч"),
    "sold_any": ("выкуплен", "выкупен", "досроч", "доссроч", "сразу"),
}

for gname, markers in groups.items():
    init = bal = val = 0
    n = 0
    for r in range(4, (ws.max_row or 0) + 1):
        addr = str(ws.cell(r, 3).value or "").strip()
        fio = str(ws.cell(r, 4).value or "").strip()
        note = str(ws.cell(r, 11).value or "").strip().casefold().replace("ё", "е")
        if not addr or not fio or addr.lower().startswith("итого"):
            continue
        if gname == "sold_any":
            ok = any(m in note for m in markers) and "действующ" not in note
        elif gname == "installment":
            ok = "действующ" in note or ( "рассроч" in note and "выкуп" not in note)
        else:
            ok = any(m in note for m in markers)
        if not ok:
            continue
        init += num(ws, r, 6)
        bal += num(ws, r, 9)
        val += num(ws, r, 10)
        n += 1
    print(gname, "n", n, "init", round(init, 2), "bal", round(bal, 2), "val", round(val, 2))

# early + immediate
init = bal = val = 0
n = 0
for r in range(4, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    note = str(ws.cell(r, 11).value or "").strip().casefold().replace("ё", "е")
    if not addr or not fio:
        continue
    if any(m in note for m in ("досроч", "доссроч", "сразу")):
        init += num(ws, r, 6)
        bal += num(ws, r, 9)
        val += num(ws, r, 10)
        n += 1
print("early+immediate", "n", n, "init", round(init, 2), "bal", round(bal, 2), "val", round(val, 2))

# installment active only from this file
init = bal = val = pay = mon = 0
n = 0
for r in range(4, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    note = str(ws.cell(r, 11).value or "").strip().casefold().replace("ё", "е")
    if not addr or not fio:
        continue
    if "действующ" in note:
        init += num(ws, r, 6)
        bal += num(ws, r, 9)
        val += num(ws, r, 10)
        n += 1
print("active installment in early file", n, round(init, 2))

wb.close()

# Installment file - only rows with monthly > 0?
p = find("рассроч")
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb["Лист1"]
init = pay = mon = 0
n = 0
for r in range(6, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    if not addr or not fio:
        continue
    if "итого" in addr.lower():
        continue
    init += num(ws, r, 6)
    pay += num(ws, r, 11)
    mon += num(ws, r, 12)
    n += 1
print("installment file ALL", n, round(init, 2), round(pay, 2), round(mon, 2))
wb.close()
