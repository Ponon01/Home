"""Excel header detection and column mapping for Budget Rent import."""

from __future__ import annotations

import re
from typing import Any

import openpyxl
from openpyxl.utils import range_boundaries

HEADER_ROW = 1
DATA_START_ROW = 2

IMPORT_SHEET_NAMES = ("2026 Март", "2026 новые март")

# All DB fields that can be filled from Excel (display order = canonical UI order)
BUDGET_RENT_FIELDS: tuple[str, ...] = (
    "valuation_object",
    "residential_complex_name",
    "street",
    "house_number",
    "apartment_number",
    "address",
    "quantity",
    "market_price",
    "proportional_market_price",
    "balance_value_2025",
    "proportional_balance_value_2025",
    "room_count",
    "total_area",
    "living_area",
    "occupied_area_sp",
    "shared_area_sp_total",
    "shared_area_per_person",
    "calculation_area_total",
    "fio",
    "cohabitation",
    "cohabitant_count",
    "note",
    "land_tax_zone_adjustment",
    "astana_land_tax_base_rate",
    "land_tax_adjustment_for_calc",
    "correction_coefficient",
    "land_tax_rate",
    "land_tax_yearly",
    "land_tax_year_ownership_adjustment",
    "amortization_monthly",
    "balance_value_month_1_2026",
    "balance_value_month_2_2026",
    "balance_value_month_3_2026",
    "balance_value_month_4_2026",
    "balance_value_month_5_2026",
    "balance_value_month_6_2026",
    "balance_value_month_7_2026",
    "balance_value_month_8_2026",
    "balance_value_month_9_2026",
    "balance_value_month_10_2026",
    "balance_value_month_11_2026",
    "balance_value_month_12_2026",
    "balance_value_month_13",
    "property_tax_year_2026",
    "total_land_and_property_tax_2026",
    "material_benefit",
    "reimbursement_cost_monthly",
    "taxable_base",
    "contract_status",
    "rental_contract_number_date",
    "attachment_note",
    "changes_note",
)

NUMERIC_FIELDS = frozenset(
    f
    for f in BUDGET_RENT_FIELDS
    if f
    not in {
        "valuation_object",
        "residential_complex_name",
        "street",
        "house_number",
        "apartment_number",
        "address",
        "fio",
        "cohabitation",
        "note",
        "contract_status",
        "rental_contract_number_date",
        "attachment_note",
        "changes_note",
    }
)

INTEGER_FIELDS = frozenset({"room_count", "cohabitant_count"})


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def header_to_field(header: str) -> str | None:
    """Map normalized Excel header text to budget_rent_records column name."""
    h = normalize_header(header)
    if not h:
        return None
    if h in {"№", "пп", "n", "no"} or h.startswith("№"):
        return None

    if "учетом принятия" in h and "земельн" in h:
        return "land_tax_year_ownership_adjustment"
    if re.search(r"балансовая\s+стоимость\s+13\s+месяц", h):
        return "balance_value_month_13"
    for month in range(12, 0, -1):
        if re.search(rf"балансовая\s+стоимость\s+{month}\s+месяц\s+2026", h):
            return f"balance_value_month_{month}_2026"
    if ("земельный" in h or "земельн" in h) and "имуществен" in h and "2026" in h:
        return "total_land_and_property_tax_2026"
    if "имущественный налог" in h and "2026" in h:
        return "property_tax_year_2026"
    if "материальная выгода" in h:
        return "material_benefit"
    if "амортизация" in h and "месяц" in h:
        return "amortization_monthly"
    if "себестоимость" in h and "возмещ" in h:
        return "reimbursement_cost_monthly"
    if "налогооблагаемая" in h or "налоооблагаемая" in h:
        return "taxable_base"
    if "статус договора" in h:
        return "contract_status"
    if "договор найма" in h:
        return "rental_contract_number_date"
    if h.startswith("приложение"):
        return "attachment_note"
    if h.startswith("изменен"):
        return "changes_note"
    if "объект оценки" in h:
        return "valuation_object"
    if "наименование жк" in h:
        return "residential_complex_name"
    if h == "улица":
        return "street"
    if h == "дом":
        return "house_number"
    if h == "квартира":
        return "apartment_number"
    if h == "адрес":
        return "address"
    if h == "количество":
        return "quantity"
    if "рыночная цена" in h and "пропорц" in h:
        return "proportional_market_price"
    if "рыночная цена" in h:
        return "market_price"
    if "балансовая стоимость" in h and "2025" in h and "пропорц" in h:
        return "proportional_balance_value_2025"
    if "балансовая стоимость" in h and "2025" in h:
        return "balance_value_2025"
    if "количество комнат" in h:
        return "room_count"
    if "общая площадь" in h and "сп" in h and "каждого" in h:
        return "shared_area_per_person"
    if "общая площадь" in h and "сп" in h:
        return "shared_area_sp_total"
    if "занимаемая площадь" in h:
        return "occupied_area_sp"
    if "итого площадь" in h:
        return "calculation_area_total"
    if "жилая площадь" in h:
        return "living_area"
    if "общая площадь" in h:
        return "total_area"
    if "фио" in h and "прожива" in h:
        return "fio"
    if "совместное проживание" in h:
        return "cohabitation"
    if "количество совместно" in h:
        return "cohabitant_count"
    if h.startswith("примечан"):
        return "note"
    if "корректировка" in h and "510" in h and "для расчета" in h:
        return "land_tax_adjustment_for_calc"
    if "корректировка" in h and "510" in h:
        return "land_tax_zone_adjustment"
    if "базовая ставка" in h and ("астана" in h or "505" in h):
        return "astana_land_tax_base_rate"
    if "кооректировочный" in h or "корректировочный" in h:
        return "correction_coefficient"
    if "ставка земельного налога" in h:
        return "land_tax_rate"
    if "земельный налог" in h and "год" in h:
        return "land_tax_yearly"

    return None


def build_merged_lookup(sheet) -> dict[tuple[int, int], Any]:
    lookup: dict[tuple[int, int], Any] = {}
    for merged_range in sheet.merged_cells.ranges:
        master = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row, col)] = master
    return lookup


def get_cell_value(sheet, row: int, col: int, merged_lookup: dict[tuple[int, int], Any]):
    if (row, col) in merged_lookup:
        return merged_lookup[(row, col)]
    return sheet.cell(row=row, column=col).value


def sheet_used_bounds(sheet) -> tuple[int, int]:
    if sheet.dimensions:
        try:
            _min_col, min_row, _max_col, max_row = range_boundaries(sheet.dimensions)
            return min_row, max_row
        except ValueError:
            pass
    return 1, sheet.max_row or 1


def build_sheet_column_map(sheet) -> tuple[dict[int, str], dict[int, str], dict[str, int]]:
    """
    Returns:
      col_to_field: Excel column index -> DB field
      col_to_header: Excel column index -> raw header label
      field_to_col: DB field -> Excel column index (first match wins)
    """
    merged_lookup = build_merged_lookup(sheet)
    col_to_field: dict[int, str] = {}
    col_to_header: dict[int, str] = {}
    field_to_col: dict[str, int] = {}

    _min_row, _max_row = sheet_used_bounds(sheet)
    max_col = sheet.max_column or 1

    for col in range(1, max_col + 1):
        raw = get_cell_value(sheet, HEADER_ROW, col, merged_lookup)
        label = str(raw).replace("\n", " ").strip() if raw is not None else ""
        if not label:
            continue
        field = header_to_field(label)
        col_to_header[col] = label
        if field is None:
            continue
        col_to_field[col] = field
        if field not in field_to_col:
            field_to_col[field] = col

    return col_to_field, col_to_header, field_to_col


def dedup_key(row: dict) -> tuple[str, str, str, str]:
    return (
        row.get("residential_complex_name") or "",
        row.get("address") or "",
        row.get("apartment_number") or "",
        row.get("fio") or "",
    )
