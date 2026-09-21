"""Excel parsing for Budget Purchase — import every column by exact header label."""

from __future__ import annotations

import re
from datetime import date, datetime
from typing import Any

from openpyxl.utils import range_boundaries


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    text = re.sub(r"\s+", " ", text).strip().lower()
    text = text.replace("стоймость", "стоимость")
    return text


def preserve_header_label(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ").replace("\r", " ")).strip()


def should_skip_header_label(label: str) -> bool:
    h = normalize_header(label)
    if not h:
        return True
    if h in {"№", "пп", "n", "no", "#"}:
        return True
    if h.startswith("№") and "п" in h:
        return True
    if re.fullmatch(r"№\s*п/?п?", h):
        return True
    return False


def build_merged_lookup(sheet) -> dict[tuple[int, int], Any]:
    lookup: dict[tuple[int, int], Any] = {}
    for merged_range in sheet.merged_cells.ranges:
        master = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row, col)] = master
    return lookup


def get_cell_value(sheet, row: int, col: int, merged_lookup: dict[tuple[int, int], Any]):
    if (row, col) in merged_lookup:
        return merged_lookup[(row, col)]
    return sheet.cell(row=row, column=col).value


def sheet_used_bounds(sheet) -> tuple[int, int, int]:
    min_row, max_row, max_col = 1, sheet.max_row or 1, sheet.max_column or 1
    if sheet.dimensions:
        try:
            _min_col, min_row, max_col, max_row = range_boundaries(sheet.dimensions)
        except ValueError:
            pass
    return min_row, max_row, max_col


def detect_header_row(sheet, merged_lookup: dict, max_scan: int = 25) -> int | None:
    _min_row, max_row, max_col = sheet_used_bounds(sheet)
    for row in range(1, min(max_scan, max_row) + 1):
        normalized: list[str] = []
        for col in range(1, max_col + 1):
            raw = get_cell_value(sheet, row, col, merged_lookup)
            if raw is not None and str(raw).strip():
                normalized.append(normalize_header(raw))
        joined = " ".join(normalized)
        if "адрес" in joined and ("фио" in joined or "договор" in joined):
            return row
    return None


def build_sheet_header_columns(
    sheet, header_row: int, merged_lookup: dict
) -> list[tuple[int, str]]:
    """All importable columns: (col_index, exact Excel header label)."""
    _min_row, _max_row, max_col = sheet_used_bounds(sheet)
    columns: list[tuple[int, str]] = []
    seen_labels: set[str] = set()
    for col in range(1, max_col + 1):
        raw = get_cell_value(sheet, header_row, col, merged_lookup)
        label = preserve_header_label(raw)
        if should_skip_header_label(label):
            continue
        if label in seen_labels:
            label = f"{label} ({col})"
        seen_labels.add(label)
        columns.append((col, label))
    return columns


def coerce_cell_value(raw: Any) -> str | int | float | None:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    if isinstance(raw, bool):
        return str(raw)
    if isinstance(raw, (int, float)):
        if isinstance(raw, float) and raw != raw:
            return None
        return raw
    s = str(raw).strip()
    if not s or s.lower() in ("nan", "null", "none"):
        return None
    s_num = s.replace("\u00a0", "").replace(" ", "").replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", s_num):
        try:
            return int(s_num) if "." not in s_num else float(s_num)
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s


def _cell_has_value(val) -> bool:
    if val is None:
        return False
    if isinstance(val, str):
        return bool(val.strip()) and val.strip().lower() not in ("nan", "null", "none")
    return True


def row_is_empty(sheet, row_idx: int, column_defs: list[tuple[int, str]], merged_lookup: dict) -> bool:
    for col, _label in column_defs:
        if _cell_has_value(get_cell_value(sheet, row_idx, col, merged_lookup)):
            return False
    return True


def extract_search_fields(excel_columns: dict[str, Any]) -> tuple[str | None, str | None, str | None]:
    address = fio = contract = None
    for label, value in excel_columns.items():
        h = normalize_header(label)
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        text = str(value).strip() if not isinstance(value, (int, float)) else str(value)
        if h == "адрес" or (h.startswith("адрес") and "жк" not in h):
            address = text
        elif "фио" in h:
            fio = text
        elif "договор купли" in h or (h.startswith("договор") and "продаж" in h and "взнос" not in h):
            if not contract or len(text) > len(contract):
                contract = text
    return address, fio, contract


def dedup_key(excel_columns: dict[str, Any]) -> tuple[str, str, str]:
    address, fio, contract = extract_search_fields(excel_columns)
    return (address or "", fio or "", contract or "")


def merge_column_order(global_order: list[str], sheet_order: list[str]) -> list[str]:
    seen = set(global_order)
    result = list(global_order)
    for label in sheet_order:
        if label not in seen:
            seen.add(label)
            result.append(label)
    return result
