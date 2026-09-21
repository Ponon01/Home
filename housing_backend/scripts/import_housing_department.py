import sys
import openpyxl
from pathlib import Path
from sqlalchemy import delete

sys.path.append(str(Path(__file__).resolve().parent.parent))

from app.db.session import SyncSessionLocal
from app.models.housing_department_record import HousingDepartmentRecord

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in ("nan", "null", "none", ""):
        return None
    return s

def clean_float(val):
    s = clean_str(val)
    if s is None:
        return None
    # Strip spaces (like "14 613 382.40") and commas
    s = s.replace(" ", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None

def clean_int(val):
    s = clean_str(val)
    if s is None:
        return None
    s = s.replace(" ", "").replace(",", "")
    # Handle floats representation like "2010.0"
    if "." in s:
        s = s.split(".")[0]
    try:
        return int(s)
    except ValueError:
        return None

def import_excel():
    data_dir = Path("data")
    target_file = data_dir / "Надо заполнить новую таблицу Кундыз.xlsx"
    excel_files = [target_file]
    
    if not target_file.exists():
        print(f"Error: Excel file not found: {target_file}")
        sys.exit(1)
        
    session = SyncSessionLocal()
    try:
        # Load existing records to prevent duplicates
        existing_rows = session.query(
            HousingDepartmentRecord.residential_complex_name,
            HousingDepartmentRecord.address,
            HousingDepartmentRecord.fio
        ).all()
        
        # Keep track of duplicates by the exact tuple
        existing_set = set()
        for r in existing_rows:
            existing_set.add((r[0], r[1], r[2]))
            
        initial_count = len(existing_set)
        print(f"Found {initial_count} existing records in database.")
        
        total_new_records = 0
        total_duplicates_skipped = 0
        
        for excel_path in excel_files:
            print(f"\nProcessing file: {excel_path.name}")
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            
            sheet_name = 'Лист2'
            if sheet_name not in wb.sheetnames:
                print(f"  Skipping file {excel_path.name}: Sheet '{sheet_name}' not found.")
                continue
                
            sheet = wb[sheet_name]
            
            # Dynamic header mapping
            col_map = {}
            for c in range(1, sheet.max_column + 1):
                v1 = sheet.cell(row=1, column=c).value
                v2 = sheet.cell(row=2, column=c).value
                h = ''
                if v1 and str(v1).strip():
                    h += str(v1).strip().lower() + ' '
                if v2 and str(v2).strip():
                    h += str(v2).strip().lower()
                h = h.strip()
                
                if h == 'жк': col_map['jk'] = c
                elif 'адрес' in h: col_map['address'] = c
                elif 'фио' in h: col_map['fio'] = c
                elif 'состав семьи' in h: col_map['family'] = c
                elif 'первоначальная стоимость' in h: col_map['initial_cost'] = c
                elif 'рыночная цена' in h: col_map['market_price'] = c
                elif 'сумма удержания' in h: col_map['reimbursement'] = c
                elif 'налогооблагаемая база' in h: col_map['taxable'] = c
                elif 'статус квартиры по приказу' in h and 'status' not in col_map: col_map['status'] = c
                elif 'количество комнат' in h: col_map['rooms'] = c
                elif 'общая площадь' in h: col_map['area'] = c
                elif 'год постройки' in h: col_map['build_year'] = c
                elif 'лицевой счет' in h: col_map['account'] = c
                elif 'период проживания' in h: col_map['period'] = c
                elif 'должность' in h: col_map['position'] = c
                elif 'подразделение' in h: col_map['department'] = c
                elif 'основание для заселения' in h: col_map['basis'] = c
                elif 'договор найма жилья' in h: col_map['rental'] = c
                elif 'договор купли продажи' in h: col_map['purchase'] = c
                elif 'график платежей' in h: col_map['schedule'] = c
                elif 'документ подтверждающий право' in h: col_map['ownership'] = c
                elif 'акт приема-передачи' in h and 'ownership' not in col_map: col_map['ownership'] = c

            print("  Detected dynamic mapping:", col_map)
            
            records_to_insert = []
            file_duplicates = 0
            
            for row_idx in range(3, sheet.max_row + 1):
                def get_val(key):
                    if key in col_map:
                        return sheet.cell(row=row_idx, column=col_map[key]).value
                    return None
                    
                residential_complex_name = clean_str(get_val('jk'))
                address = clean_str(get_val('address'))
                fio = clean_str(get_val('fio'))

                # Validation 1: row is valid only if there is at least: address, fio or residential_complex_name
                if not (address or fio or residential_complex_name):
                    continue
                    
                # Validation 2: skip footer rows, comments, service signatures
                if fio:
                    fio_lower = fio.lower()
                    if any(junk in fio_lower for junk in ["договор купли продажи", "дата", "рассрочка", "выкуплено"]):
                        continue
                        
                # Duplicate check
                key = (residential_complex_name, address, fio)
                if key in existing_set:
                    file_duplicates += 1
                    continue
                    
                existing_set.add(key)
                
                record = HousingDepartmentRecord(
                    residential_complex_name=residential_complex_name,
                    address=address,
                    fio=fio,
                    family_composition=clean_str(get_val('family')),
                    initial_cost=clean_float(get_val('initial_cost')),
                    market_price=clean_float(get_val('market_price')),
                    reimbursement_cost_monthly=clean_float(get_val('reimbursement')),
                    taxable_base=clean_float(get_val('taxable')),
                    status=clean_str(get_val('status')),
                    room_count=clean_int(get_val('rooms')),
                    total_area=clean_float(get_val('area')),
                    build_year=clean_int(get_val('build_year')),
                    personal_account=clean_str(get_val('account')),
                    residence_period=clean_str(get_val('period')),
                    position=clean_str(get_val('position')),
                    department=clean_str(get_val('department')),
                    occupancy_and_purchase_basis=clean_str(get_val('basis')),
                    rental_contract=clean_str(get_val('rental')),
                    purchase_contract=clean_str(get_val('purchase')),
                    payment_schedule=clean_str(get_val('schedule')),
                    ownership_document=clean_str(get_val('ownership')),
                )
                records_to_insert.append(record)
                
            total_duplicates_skipped += file_duplicates
            total_new_records += len(records_to_insert)
            
            if records_to_insert:
                print(f"  Inserting {len(records_to_insert)} new records from {excel_path.name}...")
                session.add_all(records_to_insert)
                session.commit()
            else:
                print(f"  No new records to insert from {excel_path.name}.")
                
            print(f"  Skipped {file_duplicates} duplicates in this file.")

        final_count = session.query(HousingDepartmentRecord).count()
        print(f"\nImport completed successfully!")
        print(f"Total new records added: {total_new_records}")
        print(f"Total duplicates skipped: {total_duplicates_skipped}")
        print(f"Final record count in DB: {final_count}")
        
    except Exception as e:
        session.rollback()
        print("Error during database import:", e)
        sys.exit(1)
    finally:
        session.close()

if __name__ == "__main__":
    import_excel()
