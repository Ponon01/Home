"""Import Budget Rent from data/Аренда.xlsx (dynamic headers, two sheets).

Sheets: "2026 Март", "2026 новые март"
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from sqlalchemy import delete, inspect, text

sys.path.append(str(Path(__file__).parent.parent))

from app.db.session import SyncSessionLocal, sync_engine
from app.db.base import Base
from app.models.budget_rent_record import BudgetRentRecord
from scripts.budget_rent_excel import (
    BUDGET_RENT_FIELDS,
    DATA_START_ROW,
    IMPORT_SHEET_NAMES,
    INTEGER_FIELDS,
    NUMERIC_FIELDS,
    build_merged_lookup,
    build_sheet_column_map,
    dedup_key,
    get_cell_value,
    sheet_used_bounds,
)


def clean_str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in ("nan", "null", "none", ""):
        return None
    return s


def clean_float(val) -> float | None:
    s = clean_str(val)
    if s is None:
        return None
    s = s.replace(" ", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def clean_int(val) -> int | None:
    s = clean_str(val)
    if s is None:
        return None
    s = s.replace(" ", "").replace(",", "")
    if "." in s:
        s = s.split(".")[0]
    try:
        return int(s)
    except ValueError:
        return None


def _cell_has_value(val) -> bool:
    if val is None:
        return False
    if isinstance(val, str):
        return bool(val.strip()) and val.strip().lower() not in ("nan", "null", "none")
    return True


def coerce_field(field: str, raw) -> object | None:
    if field in INTEGER_FIELDS:
        return clean_int(raw)
    if field in NUMERIC_FIELDS:
        return clean_float(raw)
    return clean_str(raw)


def ensure_optional_columns() -> None:
    """Add land_tax_year_ownership_adjustment if table exists without it."""
    insp = inspect(sync_engine)
    if not insp.has_table("budget_rent_records"):
        return
    cols = {c["name"] for c in insp.get_columns("budget_rent_records")}
    if "land_tax_year_ownership_adjustment" in cols:
        return
    with sync_engine.begin() as conn:
        conn.execute(
            text(
                "ALTER TABLE budget_rent_records "
                "ADD COLUMN IF NOT EXISTS land_tax_year_ownership_adjustment NUMERIC(18, 2)"
            )
        )
    print("Added column: land_tax_year_ownership_adjustment")


def print_mapping_report(sheet_name: str, col_to_field: dict[int, str], col_to_header: dict[int, str]) -> None:
    print(f"\n--- Headers & mapping: {sheet_name} ---")
    for col in sorted(col_to_header):
        field = col_to_field.get(col, "(skip)")
        print(f"  Col {col:2d}: {col_to_header[col][:70]!r} -> {field}")


def row_to_dict(
    sheet,
    row_idx: int,
    col_to_field: dict[int, str],
    merged_lookup: dict,
) -> dict:
    row: dict = {field: None for field in BUDGET_RENT_FIELDS}
    for col, field in col_to_field.items():
        raw = get_cell_value(sheet, row_idx, col, merged_lookup)
        row[field] = coerce_field(field, raw)
    return row


def row_is_completely_empty(row: dict) -> bool:
    return not any(v is not None and (not isinstance(v, str) or v.strip()) for v in row.values())


def read_sheet_rows(sheet) -> tuple[list[dict], dict[int, str], dict[int, str], int, int]:
    col_to_field, col_to_header, _field_to_col = build_sheet_column_map(sheet)
    print_mapping_report(sheet.title, col_to_field, col_to_header)

    required = ("valuation_object", "residential_complex_name", "fio", "address")
    mapped_fields = set(col_to_field.values())
    missing = [f for f in required if f not in mapped_fields]
    if missing:
        print(f"  WARNING: sheet {sheet.title!r} missing mapped fields: {missing}")

    merged_lookup = build_merged_lookup(sheet)
    _min_row, max_row = sheet_used_bounds(sheet)
    print(f"  Used range: {sheet.dimensions!r}, data rows {DATA_START_ROW}..{max_row}")
    print(f"  Merged ranges: {len(sheet.merged_cells.ranges)}")

    rows: list[dict] = []
    skipped_empty = 0
    for row_idx in range(DATA_START_ROW, max_row + 1):
        row = row_to_dict(sheet, row_idx, col_to_field, merged_lookup)
        if row_is_completely_empty(row):
            skipped_empty += 1
            continue
        rows.append(row)

    print(f"  Parsed rows: {len(rows)}, empty skipped: {skipped_empty}")
    return rows, col_to_field, col_to_header, max_row - DATA_START_ROW + 1, skipped_empty


def read_workbook_rows(target_file: Path) -> tuple[list[dict], dict[str, int], int, dict[str, dict]]:
    wb = openpyxl.load_workbook(target_file, data_only=True, read_only=False)

    missing = [n for n in IMPORT_SHEET_NAMES if n not in wb.sheetnames]
    if missing:
        wb.close()
        print(f"Error: sheet(s) not found: {missing}")
        print(f"Available: {wb.sheetnames}")
        sys.exit(1)

    seen: set[tuple[str, str, str, str]] = set()
    merged: list[dict] = []
    imported_per_sheet: dict[str, int] = {name: 0 for name in IMPORT_SHEET_NAMES}
    duplicates_skipped = 0
    sheet_meta: dict[str, dict] = {}

    for sheet_name in IMPORT_SHEET_NAMES:
        sheet = wb[sheet_name]
        sheet_rows, col_to_field, col_to_header, scanned, skipped_empty = read_sheet_rows(sheet)

        sheet_imported = 0
        sheet_duplicates = 0
        for row in sheet_rows:
            key = dedup_key(row)
            if key in seen:
                sheet_duplicates += 1
                duplicates_skipped += 1
                continue
            seen.add(key)
            merged.append(row)
            sheet_imported += 1

        imported_per_sheet[sheet_name] = sheet_imported
        sheet_meta[sheet_name] = {
            "scanned": scanned,
            "skipped_empty": skipped_empty,
            "parsed": len(sheet_rows),
            "imported": sheet_imported,
            "duplicates": sheet_duplicates,
            "col_to_field": col_to_field,
        }
        print(f"  Unique imported from sheet: {sheet_imported}, duplicates skipped: {sheet_duplicates}")

    wb.close()
    return merged, imported_per_sheet, duplicates_skipped, sheet_meta


def import_excel() -> None:
    Base.metadata.create_all(bind=sync_engine)
    ensure_optional_columns()

    target_file = Path("data") / "Аренда.xlsx"
    if not target_file.exists():
        print(f"Error: Excel file not found: {target_file.resolve()}")
        sys.exit(1)

    print("=" * 60)
    print("WHY COLUMN SHIFT HAPPENED (old import):")
    print("  Excel has columns 1=№, 2=пп before data.")
    print("  Old code used hardcoded index 2 for 'Объект оценки',")
    print("  but col 2 is 'пп' and col 3 is 'Объект оценки'.")
    print("  So 'Наименование ЖК' showed apartment type (valuation_object).")
    print("  Fix: map columns by header text dynamically per sheet.")
    print("=" * 60)
    print(f"\nFile: {target_file.resolve()}")
    print(f"Sheets: {list(IMPORT_SHEET_NAMES)}\n")

    row_dicts, per_sheet, duplicates_skipped, _meta = read_workbook_rows(target_file)

    print(f"\nDuplicates skipped (ЖК+Адрес+Квартира+ФИО): {duplicates_skipped}")
    print(f"Total unique rows: {len(row_dicts)}")

    session = SyncSessionLocal()
    try:
        deleted = session.execute(delete(BudgetRentRecord)).rowcount
        session.commit()
        print(f"Cleared budget_rent_records: {deleted}")

        records = [BudgetRentRecord(**row) for row in row_dicts]
        if records:
            session.add_all(records)
            session.commit()

        final_count = session.query(BudgetRentRecord).count()
        print("\n=== Import summary ===")
        print(f'  "2026 Март": {per_sheet["2026 Март"]} row(s)')
        print(f'  "2026 новые март": {per_sheet["2026 новые март"]} row(s)')
        print(f"  Total in Budget Rent (DB): {final_count}")

        if final_count != len(row_dicts):
            print(f"WARNING: DB {final_count} != parsed {len(row_dicts)}", file=sys.stderr)
            sys.exit(1)

        if records:
            sample = records[0]
            print("\nSample row #1 sanity check:")
            print(f"  valuation_object: {sample.valuation_object!r}")
            print(f"  residential_complex_name: {sample.residential_complex_name!r}")
            print(f"  fio: {sample.fio!r}")
    except Exception as e:
        session.rollback()
        print("Error during database import:", e)
        sys.exit(1)
    finally:
        session.close()


if __name__ == "__main__":
    import_excel()
