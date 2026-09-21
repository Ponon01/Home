# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

DATA = Path("data")


def find(*keys):
    for p in sorted(DATA.glob("*.xlsx")):
        if p.name.startswith("~$") or p.name.startswith("_"):
            continue
        low = p.name.casefold().replace("ё", "е")
        if all(k in low for k in keys):
            return p
    return None


def num(ws, r, c):
    v = ws.cell(r, c).value
    try:
        return float(v or 0)
    except Exception:
        return 0


p = find("доссроч", "выкуп") or find("досроч", "выкуп")
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
markers = ("выкуплен", "выкупен", "досроч", "доссроч", "сразу", "100%")
init = bal = val = 0
n = 0
for r in range(4, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    note = str(ws.cell(r, 11).value or "").strip()
    if not addr or not fio or addr.lower().startswith("итого"):
        continue
    low = note.casefold().replace("ё", "е")
    if not any(x in low for x in markers):
        continue
    init += num(ws, r, 6)
    bal += num(ws, r, 9)
    val += num(ws, r, 10)
    n += 1
print("early markers n", n, "init", round(init, 2), "bal", round(bal, 2), "val", round(val, 2))
wb.close()

# fund totals
p = find("информац", "фонд")
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb["Свод по ЖК русс"]
tot = nfs = fs = rent = guest = 0
for r in range(8, 24):
    name = str(ws.cell(r, 1).value or "")
    if not name or name.upper().startswith("ИТОГО"):
        continue

    def ni(c):
        v = ws.cell(r, c).value
        try:
            return int(v or 0)
        except Exception:
            return 0

    tot += ni(3)
    nfs += ni(4)
    fs += ni(5)
    rent += ni(17)
    guest += ni(18)
print("fund", tot, nfs, fs, rent, guest)
wb.close()
