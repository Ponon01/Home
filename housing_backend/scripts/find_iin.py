import openpyxl
from pathlib import Path

def find_iin(filename):
    path = Path("data") / filename
    if not path.exists():
        print(f"File {filename} not found")
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        for r_idx, row in enumerate(ws.iter_rows(max_row=10)):
            row_vals = [cell.value for cell in row]
            for col_idx, val in enumerate(row_vals):
                if val and "ИИН" in str(val).upper():
                    print(f"[{filename}] Found IIN in Sheet '{name}', Row {r_idx + 1}, Column {col_idx + 1}: {val}")
    wb.close()

if __name__ == "__main__":
    find_iin("1Расчет по квартирам 2026 (АРЕНДА).xlsx")
    find_iin("рассрочка и выкупленные общее поступление 2222.xlsx")
