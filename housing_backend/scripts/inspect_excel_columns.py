import openpyxl
from pathlib import Path

def inspect_file(filename, out_file):
    path = Path("data") / filename
    if not path.exists():
        out_file.write(f"File {filename} not found\n")
        return
    out_file.write(f"\n=== {filename} ===\n")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        out_file.write(f"Sheet: '{name}'\n")
        # Print first 10 rows to see structure
        for r_idx, row in enumerate(ws.iter_rows(max_row=10)):
            row_vals = [cell.value for cell in row]
            if any(row_vals is not None for row_vals in row_vals):
                # Clean up values to be printable strings
                str_vals = [str(v) if v is not None else "" for v in row_vals]
                out_file.write(f"  Row {r_idx + 1}: {str_vals[:20]}\n")
    wb.close()

if __name__ == "__main__":
    out_path = Path("c:/Users/Admin/Desktop/housing_backend/scratch_inspect_columns.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        inspect_file("1Расчет по квартирам 2026 (АРЕНДА).xlsx", f)
        inspect_file("рассрочка и выкупленные общее поступление 2222.xlsx", f)
    print("Done inspecting. Output saved to scratch_inspect_columns.txt")
