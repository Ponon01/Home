# -*- coding: utf-8 -*-
"""Подбор подмножества early-строк под контрольные bal/val."""
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

DATA = Path("data")


def find(*keys):
    for p in sorted(DATA.glob("*.xlsx")):
        if p.name.startswith("~$") or p.name.startswith("_"):
            continue
        low = p.name.casefold().replace("ё", "е")
        if all(k in low for k in keys):
            return p
    return None


def num(ws, r, c):
    try:
        return float(ws.cell(r, c).value or 0)
    except Exception:
        return 0


p = find("доссроч", "выкуп") or find("досроч", "выкуп")
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
markers = ("выкуплен", "выкупен", "досроч", "доссроч", "сразу", "100%")
rows = []
for r in range(4, (ws.max_row or 0) + 1):
    addr = str(ws.cell(r, 3).value or "").strip()
    fio = str(ws.cell(r, 4).value or "").strip()
    note = str(ws.cell(r, 11).value or "").strip()
    if not addr or not fio or addr.lower().startswith("итого"):
        continue
    low = note.casefold().replace("ё", "е")
    if not any(x in low for x in markers):
        continue
    rows.append(
        {
            "r": r,
            "note": note,
            "init": num(ws, r, 6),
            "bal": num(ws, r, 9),
            "val": num(ws, r, 10),
        }
    )
wb.close()

target_init = 652289417.0
target_bal = 604881499.99
target_val = 387364199.81
print("n", len(rows), "sum", sum(x["init"] for x in rows), sum(x["bal"] for x in rows), sum(x["val"] for x in rows))
print("diff bal", sum(x["bal"] for x in rows) - target_bal)
print("diff val", sum(x["val"] for x in rows) - target_val)

# find single row that explains val diff
dval = sum(x["val"] for x in rows) - target_val
dbal = sum(x["bal"] for x in rows) - target_bal
print("looking for rows near dval", dval, "dbal", dbal)
for x in rows:
    if abs(x["val"] - dval) < 1 or abs(x["bal"] - dbal) < 1:
        print("match1", x)
for a, b in combinations(rows, 2):
    if abs(a["val"] + b["val"] - dval) < 1 and abs(a["bal"] + b["bal"] - dbal) < 1:
        print("match2", a, b)
        break

# group by note
from collections import defaultdict

g = defaultdict(lambda: [0, 0.0, 0.0, 0.0])
for x in rows:
    k = x["note"]
    g[k][0] += 1
    g[k][1] += x["init"]
    g[k][2] += x["bal"]
    g[k][3] += x["val"]
for k, v in g.items():
    print("group", v[0], k, round(v[1], 2), round(v[2], 2), round(v[3], 2))
