import asyncio
import sys
import os
import openpyxl
from collections import defaultdict
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from sqlalchemy import select
from app.models.apartment import Apartment

DATABASE_URL = "postgresql+asyncpg://postgres:postgres@localhost:5432/astana_opera_housing"

def normalize_name(name):
    if not name:
        return ""
    name = name.lower().strip()
    name = name.replace(" ", "").replace("-", "").replace("—", "")
    if "лазурный" in name:
        return "лазурныйквартал"
    if "сапа" in name or "сити" in name:
        return "сапа2007"
    if "азирбаев" in name:
        return "улкаазирбаева"
    return name

async def main():
    engine = create_async_engine(DATABASE_URL, echo=False)
    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    
    # 1. Read and parse Excel data
    print("Loading Книга1.xlsx...")
    wb = openpyxl.load_workbook("data/Книга1.xlsx", data_only=True)
    sheet = wb.worksheets[0]
    print(f"Sheet loaded: {sheet.title}")
    
    excel_apts = defaultdict(list)
    for r in range(2, sheet.max_row + 1):
        idx = sheet.cell(row=r, column=2).value
        apt_num = sheet.cell(row=r, column=7).value
        full_addr = sheet.cell(row=r, column=8).value
        
        if not full_addr:
            continue
            
        complex_guess = ""
        full_addr_lower = full_addr.lower()
        if "лазурный" in full_addr_lower:
            complex_guess = "Лазурный квартал"
        elif "зерде" in full_addr_lower:
            complex_guess = "Зерде"
        elif "сармат" in full_addr_lower:
            complex_guess = "Сармат"
        elif "хан-тенгри" in full_addr_lower or "хан тенгри" in full_addr_lower:
            complex_guess = "Хан-тенгри"
        elif "виктория" in full_addr_lower:
            complex_guess = "Виктория"
        elif "жагалау" in full_addr_lower:
            complex_guess = "Жагалау-3"
        elif "сапа" in full_addr_lower:
            complex_guess = "Сапа-2007"
        elif "акку" in full_addr_lower:
            complex_guess = "Акку"
        elif "общежитие" in full_addr_lower:
            complex_guess = "Общежитие"
        elif "азирбаев" in full_addr_lower:
            complex_guess = "ул. К. Азирбаева"
        elif "браво" in full_addr_lower:
            complex_guess = "Браво"
        elif "compass" in full_addr_lower:
            complex_guess = "Compass North"
        elif "respublika" in full_addr_lower or "республика" in full_addr_lower:
            if "республики 81" not in full_addr_lower:
                complex_guess = "Respublika"
            else:
                complex_guess = "Общежитие"
        elif "нур-сая" in full_addr_lower or "нур сая" in full_addr_lower or "нурсая" in full_addr_lower:
            complex_guess = "НУР-САЯ"
        elif "москва" in full_addr_lower:
            complex_guess = "Москва"
            
        if not complex_guess:
            continue
            
        amort = sheet.cell(row=r, column=47).value
        deduct = sheet.cell(row=r, column=48).value
        tax = sheet.cell(row=r, column=49).value
        
        try:
            amort = float(amort) if amort is not None else 0.0
        except:
            amort = 0.0
        try:
            deduct = float(deduct) if deduct is not None else 0.0
        except:
            deduct = 0.0
        try:
            tax = float(tax) if tax is not None else 0.0
        except:
            tax = 0.0
            
        apt_str = str(apt_num).strip() if apt_num is not None else ""
        if not apt_str and "-" in full_addr:
            parts = full_addr.split("-")
            if parts:
                apt_str = parts[-1].strip()
                
        key = (normalize_name(complex_guess), apt_str)
        excel_apts[key].append({
            "amortization": amort,
            "deduction": deduct,
            "tax": tax
        })
        
    print(f"Total parsed Excel apartments/groups: {len(excel_apts)}")
    
    # 2. Update DB apartments
    async with async_session() as session:
        res = await session.execute(select(Apartment))
        db_apts = res.scalars().all()
        
        updated_count = 0
        for a in db_apts:
            c_norm = normalize_name(a.residential_complex_name)
            apt_str = str(a.apartment_number).strip() if a.apartment_number else ""
            key = (c_norm, apt_str)
            
            if key in excel_apts:
                rows = excel_apts[key]
                # Sum values for the apartment
                total_amort = sum(r["amortization"] for r in rows)
                total_deduct = sum(r["deduction"] for r in rows)
                total_tax = sum(r["tax"] for r in rows)
                
                a.amortization_cost = total_amort
                a.monthly_deduction = total_deduct
                a.taxable_base = total_tax
                updated_count += 1
                
                if updated_count <= 5:
                    print(f"Sample updated: {a.residential_complex_name} Apt {a.apartment_number} -> "
                          f"amort: {a.amortization_cost}, deduct: {a.monthly_deduction}, tax: {a.taxable_base}")
            else:
                # If not matched, we can set to None or 0 or keep as is.
                # Let's set to None or 0 to be clean, or keep as None since columns are optional.
                pass
                
        await session.commit()
        print(f"Successfully updated {updated_count} apartments in database.")

if __name__ == "__main__":
    asyncio.run(main())
