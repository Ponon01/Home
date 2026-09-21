import openpyxl
from pathlib import Path

def inspect_file(filename, sheets=None, header_row=1):
    path = Path("data") / filename
    if not path.exists():
        print(f"File {filename} not found")
        return
    print(f"\n=== {filename} ===")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    target_sheets = sheets if sheets else wb.sheetnames
    for name in target_sheets:
        if name not in wb.sheetnames:
            print(f"Sheet {name} not found")
            continue
        ws = wb[name]
        row_vals = [cell.value for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
        print(f"Sheet '{name}' (header row {header_row}):")
        for i, val in enumerate(row_vals):
            print(f"  [{i}] {val}")
    wb.close()

if __name__ == "__main__":
    inspect_file("Аренда.xlsx", ["2026 Март", "2026 новые март"], header_row=1)
    inspect_file("рассрочка и выкупленные общее поступление.xlsx", header_row=4)
