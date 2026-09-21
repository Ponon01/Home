import openpyxl

wb = openpyxl.load_workbook("data/Книга1.xlsx", data_only=True)
sheet = wb.worksheets[0]

with open("scratch_excel_rows.txt", "w", encoding="utf-8") as f:
    f.write("Row | Index | Rooms | House | Apt | FullAddress | Amortization (AU) | Deduction (AV) | TaxBase (AW)\n")
    for r in range(2, 200):
        idx = sheet.cell(row=r, column=2).value
        rooms = sheet.cell(row=r, column=3).value
        house = sheet.cell(row=r, column=6).value
        apt = sheet.cell(row=r, column=7).value
        full_addr = sheet.cell(row=r, column=8).value
        amort = sheet.cell(row=r, column=47).value
        deduct = sheet.cell(row=r, column=48).value
        tax = sheet.cell(row=r, column=49).value
        
        if idx is not None or full_addr is not None:
            f.write(f"{r} | {idx} | {rooms} | {house} | {apt} | {full_addr} | {amort} | {deduct} | {tax}\n")
print("Done writing Excel dump.")
